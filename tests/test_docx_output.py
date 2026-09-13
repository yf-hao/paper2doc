from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from httpx2 import Request, Response
from openai import InternalServerError

from paper2doc.docx_writer import write_docx
from paper2doc.models import Paragraph


class IdentityTranslator:
    def translate(self, text):
        return f"中文：{text}"


class BatchTranslator:
    def __init__(self):
        self.batches = []

    def translate_batch(self, units):
        self.batches.append(units)
        return {unit_id: f"中文：{text}" for unit_id, text in units}


class SplitFallbackTranslator:
    def __init__(self):
        self.calls = []

    def translate_batch(self, units):
        self.calls.append(len(units))
        if len(units) > 1:
            raise RuntimeError("invalid batch response")
        return {unit_id: f"中文：{text}" for unit_id, text in units}

    def translate(self, text):
        return f"中文：{text}"


class RetryThenSuccessTranslator:
    def __init__(self):
        self.attempts = 0

    def translate_batch(self, units):
        self.attempts += 1
        if self.attempts < 3:
            raise RuntimeError("temporary malformed response")
        return {unit_id: f"中文：{text}" for unit_id, text in units}


class BadGatewayThenSuccessTranslator:
    def __init__(self):
        self.attempts = 0

    def translate_batch(self, units):
        self.attempts += 1
        if self.attempts < 3:
            raise InternalServerError(
                "bad gateway",
                response=Response(
                    502,
                    request=Request("POST", "https://example.test/v1/chat/completions"),
                ),
                body=None,
            )
        return {unit_id: f"中文：{text}" for unit_id, text in units}


class RecordingProgress:
    def __init__(self):
        self.started = None
        self.stages = []
        self.updates = []
        self.finished = False
        self.failed = False

    def stage(self, message):
        self.stages.append(message)

    def start(self, total, description):
        self.started = (total, description)

    def update(self, completed):
        self.updates.append(completed)

    def finish(self):
        self.finished = True

    def fail(self):
        self.failed = True


def test_docx_has_separate_english_and_chinese_paragraphs(tmp_path):
    output = tmp_path / "result.docx"
    write_docx(
        [Paragraph(1, 0, (0, 0, 100, 20), "An English paragraph.")],
        [],
        IdentityTranslator(),
        output,
    )
    document = Document(output)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "An English paragraph.",
    ]
    assert len(document.tables) == 1
    translation_cell = document.tables[0].cell(0, 0)
    assert translation_cell.text == "中文：An English paragraph."
    assert document.tables[0].style.name == "Table Grid"
    assert translation_cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    borders = translation_cell._tc.tcPr.find(qn("w:tcBorders"))
    assert borders is not None
    assert {border.get(qn("w:val")) for border in borders} == {"nil"}
    shading = translation_cell._tc.tcPr.find(qn("w:shd"))
    assert shading is not None
    assert shading.get(qn("w:fill")) == "B7D4EF"
    assert output.is_file()


def test_progress_tracks_successful_translations(tmp_path):
    output = tmp_path / "result.docx"
    progress = RecordingProgress()
    write_docx(
        [Paragraph(1, 0, (0, 0, 100, 20), "An English paragraph.")],
        [],
        IdentityTranslator(),
        output,
        progress=progress,
    )
    assert progress.started == (1, "Translating paragraphs")
    assert progress.updates == [1]
    assert progress.finished is True
    assert progress.failed is False


def test_docx_writer_batches_translations_by_character_limit(tmp_path):
    output = tmp_path / "result.docx"
    translator = BatchTranslator()
    progress = RecordingProgress()
    paragraphs = [
        Paragraph(1, 0, (0, 0, 100, 20), "12345"),
        Paragraph(2, 0, (0, 30, 100, 50), "67890"),
        Paragraph(3, 0, (0, 60, 100, 80), "abc"),
    ]

    write_docx(
        paragraphs,
        [],
        translator,
        output,
        progress=progress,
        batch_size=10,
        batch_min_size=1,
    )

    assert translator.batches == [
        [("paragraph:1", "12345"), ("paragraph:2", "67890")],
        [("paragraph:3", "abc")],
    ]
    assert progress.stages == [
        "Translating batch 1/2 (10 characters)... waiting for response",
        "Translated batch 1/2: 2/3 paragraphs",
        "Translating batch 2/2 (3 characters)... waiting for response",
        "Translated batch 2/2: 3/3 paragraphs",
    ]
    assert progress.updates == [1, 2, 3]
    document = Document(output)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "12345",
        "67890",
        "abc",
    ]
    assert [table.cell(0, 0).text for table in document.tables] == [
        "中文：12345",
        "中文：67890",
        "中文：abc",
    ]


def test_batch_progress_marks_batch_below_recommended_minimum(tmp_path):
    output = tmp_path / "result.docx"
    translator = BatchTranslator()
    progress = RecordingProgress()
    paragraphs = [
        Paragraph(1, 0, (0, 0, 100, 20), "12345"),
        Paragraph(2, 0, (0, 30, 100, 50), "67890"),
        Paragraph(3, 0, (0, 60, 100, 80), "abc"),
    ]

    write_docx(
        paragraphs,
        [],
        translator,
        output,
        progress=progress,
        batch_size=10,
        batch_min_size=6,
    )

    assert progress.stages[-1] == "Translated batch 2/2: 3/3 paragraphs (below minimum)"


def test_invalid_batch_response_falls_back_to_smaller_batches(tmp_path):
    output = tmp_path / "result.docx"
    translator = SplitFallbackTranslator()
    paragraphs = [
        Paragraph(1, 0, (0, 0, 100, 20), "12345"),
        Paragraph(2, 0, (0, 30, 100, 50), "67890"),
        Paragraph(3, 0, (0, 60, 100, 80), "abc"),
    ]

    write_docx(
        paragraphs,
        [],
        translator,
        output,
        batch_size=10,
        batch_min_size=1,
    )

    assert translator.calls == [2, 2, 2, 1, 1, 1]


def test_batch_retries_twice_before_succeeding(tmp_path):
    output = tmp_path / "result.docx"
    translator = RetryThenSuccessTranslator()

    write_docx(
        [Paragraph(1, 0, (0, 0, 100, 20), "12345")],
        [],
        translator,
        output,
        batch_size=10,
        batch_min_size=1,
    )

    assert translator.attempts == 3


def test_batch_retries_http_502_twice_before_succeeding(tmp_path):
    output = tmp_path / "result.docx"
    translator = BadGatewayThenSuccessTranslator()

    write_docx(
        [Paragraph(1, 0, (0, 0, 100, 20), "12345")],
        [],
        translator,
        output,
        batch_size=10,
        batch_min_size=1,
    )

    assert translator.attempts == 3


def test_translation_failure_is_not_silently_replaced(tmp_path):
    class FailingTranslator:
        def translate(self, text):
            raise RuntimeError("translation unavailable")

    output = tmp_path / "result.docx"
    progress = RecordingProgress()
    with pytest.raises(RuntimeError, match="translation unavailable"):
        write_docx(
            [Paragraph(1, 0, (0, 0, 100, 20), "An English paragraph.")],
            [],
            FailingTranslator(),
            output,
            progress=progress,
        )
    assert progress.updates == []
    assert progress.finished is False
    assert progress.failed is True
