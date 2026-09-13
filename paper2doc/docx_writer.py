from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from openai import APIStatusError

from .models import ImageBlock, Paragraph, TableBlock
from .layout_analyzer import reading_order
from .progress import NullProgress, ProgressReporter


FONT_CHINESE = "宋体"
FONT_ENGLISH = "Times New Roman"
BODY_SIZE_PT = 10.5
PAGE_MARGIN_CM = 2.54
FIRST_LINE_INDENT_CM = 0.74
TABLE_WIDTH_CM = 21 - (PAGE_MARGIN_CM * 2)
RETRYABLE_BATCH_STATUS_CODES = frozenset({408, 429, 500, 502, 503, 504})


def _format(paragraph, font_name: str, after: float, indent: bool = True):
    paragraph.paragraph_format.first_line_indent = Cm(FIRST_LINE_INDENT_CM) if indent else Cm(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(BODY_SIZE_PT)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
            font_name,
        )


def _add_text(document, text: str, font_name: str, after: float):
    paragraph = document.add_paragraph(text)
    _format(paragraph, font_name, after)
    return paragraph


def _add_translation_table(document, text: str):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(TABLE_WIDTH_CM)
    cell = table.cell(0, 0)
    cell.width = Cm(TABLE_WIDTH_CM)
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    cell_properties.append(borders)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "B7D4EF")
    cell_properties.append(shading)
    paragraph = cell.paragraphs[0]
    paragraph.text = text
    _format(paragraph, FONT_CHINESE, 4, indent=False)
    return table


def _set_table_borders(table, vertical: bool = True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{side}"))
        if border is None:
            border = OxmlElement(f"w:{side}")
            borders.append(border)
        border.set(qn("w:val"), "single" if vertical or side in {"top", "bottom", "insideH"} else "nil")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "808080")


def _set_header_repeat(cell):
    tr_pr = cell._tc.getparent().get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _add_native_table(document, table_block: TableBlock, translations: dict[str, str], chinese: bool = False):
    table = document.add_table(rows=max(table_block.rows, 1), cols=max(table_block.columns, 1))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total_width = max(table_block.bbox[2] - table_block.bbox[0], 1)
    widths = [total_width / max(table_block.columns, 1)] * max(table_block.columns, 1)
    for cell in table_block.cells:
        if cell.column < len(widths):
            widths[cell.column] = max(widths[cell.column], cell.bbox[2] - cell.bbox[0])
    width_total = sum(widths) or 1
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(TABLE_WIDTH_CM * width / width_total)

    cells = sorted(table_block.cells, key=lambda item: (item.row, item.column))
    occupied: set[tuple[int, int]] = set()
    for source_cell in cells:
        if source_cell.row >= table_block.rows or source_cell.column >= table_block.columns:
            continue
        if (source_cell.row, source_cell.column) in occupied:
            continue
        target = table.cell(source_cell.row, source_cell.column)
        if source_cell.row_span > 1 or source_cell.col_span > 1:
            end_row = min(table_block.rows - 1, source_cell.row + source_cell.row_span - 1)
            end_col = min(table_block.columns - 1, source_cell.column + source_cell.col_span - 1)
            target = target.merge(table.cell(end_row, end_col))
            occupied.update(
                (row, column)
                for row in range(source_cell.row, end_row + 1)
                for column in range(source_cell.column, end_col + 1)
            )
        text = source_cell.text
        if chinese and text:
            text = translations.get(f"table:{table_block.id}:r{source_cell.row}:c{source_cell.column}", "")
        target.text = text
        target.width = Cm(TABLE_WIDTH_CM * widths[min(source_cell.column, len(widths) - 1)] / width_total)
        paragraph = target.paragraphs[0]
        _format(paragraph, FONT_CHINESE if chinese else FONT_ENGLISH, 2, indent=False)
        if source_cell.is_header:
            for run in paragraph.runs:
                run.bold = True
            _set_header_repeat(target)
    _set_table_borders(table, vertical=table_block.vertical_rules)
    return table


def _translate_one(translator, text: str) -> str:
    value = translator.translate(text) if hasattr(translator, "translate") else translator(text)
    if not value or not value.strip():
        raise RuntimeError("Translation returned empty text")
    return value.strip()


def _chunk_translation_units(
    units: list[tuple[str, str]],
    max_chars: int,
):
    if max_chars <= 0:
        raise ValueError("batch_size must be greater than zero")
    batch: list[tuple[str, str]] = []
    batch_chars = 0
    for unit in units:
        text_chars = len(unit[1])
        if batch and batch_chars + text_chars > max_chars:
            yield batch
            batch = []
            batch_chars = 0
        batch.append(unit)
        batch_chars += text_chars
    if batch:
        yield batch


def _translation_units(
    paragraphs: list[Paragraph],
    attached: dict[int | None, list[ImageBlock]],
    tables: list[TableBlock] | None = None,
) -> list[tuple[str, str]]:
    units: list[tuple[str, str]] = []
    for paragraph in paragraphs:
        if paragraph.is_caption:
            continue
        units.append((f"paragraph:{paragraph.id}", paragraph.text))
        for image in sorted(attached.get(paragraph.id, []), key=lambda item: (item.page, item.bbox[1])):
            if image.caption:
                units.append((f"image-caption:{image.id}", image.caption))

    for image in sorted(attached.get(None, []), key=lambda item: (item.page, item.bbox[1])):
        if image.caption:
            units.append((f"image-caption:{image.id}", image.caption))

    attached_captions = {image.caption for images in attached.values() for image in images if image.caption}
    for paragraph in paragraphs:
        if paragraph.is_caption and paragraph.text not in attached_captions:
            units.append((f"paragraph-caption:{paragraph.id}", paragraph.text))
    for table in tables or []:
        if table.caption:
            units.append((f"table:{table.id}:caption", table.caption))
        for cell in table.cells:
            if cell.text.strip():
                units.append((f"table:{table.id}:r{cell.row}:c{cell.column}", cell.text))
    return units


def _translate_units(
    translator,
    units: list[tuple[str, str]],
    batch_size: int,
    batch_min_size: int,
    progress: ProgressReporter,
) -> dict[str, str]:
    if batch_min_size <= 0:
        raise ValueError("batch_min_size must be greater than zero")
    if batch_min_size > batch_size:
        raise ValueError("batch_min_size cannot be greater than batch_size")
    translations: dict[str, str] = {}
    completed = 0
    batch_translator = getattr(translator, "translate_batch", None)
    if callable(batch_translator):
        batches = list(_chunk_translation_units(units, batch_size))
        total_batches = len(batches)
        for batch_number, batch in enumerate(batches, start=1):
            batch_chars = sum(len(text) for _unit_id, text in batch)
            short_batch_suffix = " (below minimum)" if batch_chars < batch_min_size else ""
            progress.stage(
                f"Translating batch {batch_number}/{total_batches} "
                f"({batch_chars} characters)... waiting for response"
            )
            batch_translations = _translate_batch_with_fallback(
                translator,
                batch,
                progress,
                f"{batch_number}/{total_batches}",
            )
            for unit_id, _text in batch:
                translations[unit_id] = batch_translations[unit_id]
                completed += 1
                progress.update(completed)
            progress.stage(
                f"Translated batch {batch_number}/{total_batches}: "
                f"{completed}/{len(units)} paragraphs{short_batch_suffix}"
            )
        return translations

    for unit_id, text in units:
        translations[unit_id] = _translate_one(translator, text)
        completed += 1
        progress.update(completed)
    return translations


def _validate_batch_translations(
    batch: list[tuple[str, str]],
    batch_translations,
) -> dict[str, str]:
    if not isinstance(batch_translations, dict):
        raise RuntimeError("Batch translator must return a dictionary keyed by paragraph ID")
    expected_ids = {unit_id for unit_id, _text in batch}
    returned_ids = set(batch_translations)
    if returned_ids != expected_ids:
        missing_ids = ", ".join(sorted(expected_ids - returned_ids))
        unknown_ids = ", ".join(sorted(returned_ids - expected_ids))
        details = []
        if missing_ids:
            details.append(f"missing: {missing_ids}")
        if unknown_ids:
            details.append(f"unknown: {unknown_ids}")
        raise RuntimeError(f"Batch translation markers do not match input ({'; '.join(details)})")
    validated: dict[str, str] = {}
    for unit_id, _text in batch:
        value = batch_translations[unit_id]
        if not isinstance(value, str) or not value.strip():
            raise RuntimeError(f"Batch translation returned empty text for paragraph marker: {unit_id}")
        validated[unit_id] = value.strip()
    return validated


def _translate_batch_with_fallback(
    translator,
    batch: list[tuple[str, str]],
    progress: ProgressReporter,
    batch_label: str,
) -> dict[str, str]:
    total_attempts = 3
    for attempt in range(1, total_attempts + 1):
        try:
            return _validate_batch_translations(batch, translator.translate_batch(batch))
        except RuntimeError:
            error_suffix = ""
        except APIStatusError as exc:
            if exc.status_code not in RETRYABLE_BATCH_STATUS_CODES:
                raise
            error_suffix = f"; HTTP {exc.status_code}"
        if attempt < total_attempts:
            progress.stage(
                f"Batch {batch_label} failed (attempt {attempt}/{total_attempts}"
                f"{error_suffix}); retrying"
            )

    if len(batch) == 1:
        unit_id, text = batch[0]
        progress.stage(
            f"Batch {batch_label} failed after 2 retries for {unit_id}; "
            "retrying as a single paragraph"
        )
        return {unit_id: _translate_one(translator, text)}

    midpoint = len(batch) // 2
    progress.stage(
        f"Batch {batch_label} failed after 2 retries; downgrading to "
        f"{len(batch[:midpoint])} and {len(batch[midpoint:])} paragraphs"
    )
    left = _translate_batch_with_fallback(translator, batch[:midpoint], progress, batch_label)
    right = _translate_batch_with_fallback(translator, batch[midpoint:], progress, batch_label)
    return {**left, **right}


def write_docx(
    paragraphs: list[Paragraph],
    images: list[ImageBlock],
    translator,
    output_path: str | Path,
    progress: ProgressReporter | None = None,
    translation_description: str = "Translating paragraphs",
    batch_size: int = 7000,
    batch_min_size: int = 6000,
    tables: list[TableBlock] | None = None,
) -> Path:
    progress = progress or NullProgress()
    tables = list(tables or [])
    attached = {image.parent_paragraph_id: [] for image in images}
    for image in images:
        attached.setdefault(image.parent_paragraph_id, []).append(image)

    units = _translation_units(paragraphs, attached, tables)
    progress.start(len(units), translation_description)
    try:
        translations = _translate_units(translator, units, batch_size, batch_min_size, progress)
    except Exception:
        progress.fail()
        raise
    progress.finish()

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(PAGE_MARGIN_CM)
    section.bottom_margin = Cm(PAGE_MARGIN_CM)
    section.left_margin = Cm(PAGE_MARGIN_CM)
    section.right_margin = Cm(PAGE_MARGIN_CM)

    elements = reading_order([*paragraphs, *tables, *images])
    rendered_images: set[int] = set()
    for element in elements:
        if isinstance(element, Paragraph):
            if element.is_caption:
                continue
            _add_text(document, element.text, FONT_ENGLISH, 2)
            _add_translation_table(document, translations[f"paragraph:{element.id}"])
            for image in sorted(attached.get(element.id, []), key=lambda item: (item.page, item.bbox[1])):
                image_paragraph = document.add_paragraph()
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = image_paragraph.add_run()
                width = min(max(image.bbox[2] - image.bbox[0], 1), 451)
                run.add_picture(BytesIO(image.image_bytes), width=Pt(width))
                image_paragraph.paragraph_format.space_after = Pt(2)
                rendered_images.add(image.id)
                if image.caption:
                    _add_text(document, image.caption, FONT_ENGLISH, 1)
                    _add_translation_table(document, translations[f"image-caption:{image.id}"])
        elif isinstance(element, TableBlock):
            if element.caption:
                _add_text(document, element.caption, FONT_ENGLISH, 1)
                _add_translation_table(document, translations[f"table:{element.id}:caption"])
            if element.cells and element.rows and element.columns:
                _add_native_table(document, element, translations)
                _add_native_table(document, element, translations, chinese=True)
            elif element.fallback_image_bytes:
                image_paragraph = document.add_paragraph()
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_paragraph.add_run().add_picture(
                    BytesIO(element.fallback_image_bytes),
                    width=Pt(min(max(element.bbox[2] - element.bbox[0], 1), 451)),
                )

    # Preserve images that could not be associated instead of silently dropping them.
    for image in sorted(attached.get(None, []), key=lambda item: (item.page, item.bbox[1])):
        if image.id in rendered_images:
            continue
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(
            BytesIO(image.image_bytes), width=Pt(min(max(image.bbox[2] - image.bbox[0], 1), 451))
        )
        if image.caption:
            _add_text(document, image.caption, FONT_ENGLISH, 1)
            _add_translation_table(document, translations[f"image-caption:{image.id}"])

    attached_captions = {image.caption for images in attached.values() for image in images if image.caption}
    for paragraph in paragraphs:
        if paragraph.is_caption and paragraph.text not in attached_captions:
            _add_text(document, paragraph.text, FONT_ENGLISH, 1)
            _add_translation_table(document, translations[f"paragraph-caption:{paragraph.id}"])

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=False, exist_ok=True) if output_path.parent != Path(".") else None
    document.save(output_path)
    return output_path
